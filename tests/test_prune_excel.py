
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from ops.prune_excel import prune_sheet


def test_prune_keeps_latest_distinct_dates_and_all_rows_for_each_date():
    wb = Workbook()
    ws = wb.active
    ws.title = "sectors"
    ws.append(["sector_id", "date", "sector_name"])

    start = date(2026, 1, 1)
    sector_id = 1

    for offset in range(20):
        day = start + timedelta(days=offset)
        ws.append([sector_id, day.isoformat(), "Nifty Bank"])
        sector_id += 1

        # Both rows for the latest date must survive.
        if offset == 19:
            ws.append([sector_id, day.isoformat(), "Nifty IT"])

    removed = prune_sheet(ws, retention_limit=15)

    retained_dates = [
        ws.cell(row=r, column=2).value
        for r in range(2, ws.max_row + 1)
    ]

    assert removed == 5
    assert len(set(retained_dates)) == 15
    assert all(value >= "2026-01-06" for value in retained_dates)
    assert retained_dates.count("2026-01-20") == 2


def test_prune_normalizes_header_and_preserves_undated_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Market_Vibe"
    ws.append(["Date ", "value"])

    for offset in range(17):
        day = date(2026, 1, 1) + timedelta(days=offset)
        ws.append([day.isoformat(), offset])

    ws.append([None, "manual note without a date"])

    removed = prune_sheet(ws, retention_limit=15)

    assert removed == 2
    assert any(
        ws.cell(r, 2).value == "manual note without a date"
        for r in range(2, ws.max_row + 1)
    )


def test_prune_updates_table_range_after_row_deletions():
    wb = Workbook()
    ws = wb.active
    ws.title = "DailyData"
    ws.append(["date", "value"])

    for offset in range(20):
        day = date(2026, 1, 1) + timedelta(days=offset)
        ws.append([day.isoformat(), offset])

    table = Table(displayName="DailyDataTable", ref="A1:B21")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
    )
    ws.add_table(table)

    removed = prune_sheet(ws, retention_limit=15)

    assert removed == 5
    assert ws.tables["DailyDataTable"].ref == "A1:B16"


def test_sheet_without_date_column_is_unchanged():
    wb = Workbook()
    ws = wb.active
    ws.title = "Metadata"
    ws.append(["key", "value"])
    ws.append(["source", "manual"])

    removed = prune_sheet(ws, retention_limit=15)

    assert removed == 0
    assert ws.max_row == 2
    assert ws["B2"].value == "manual"
