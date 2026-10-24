"""Builders for small, throw-away workbooks/databases used by the tests."""

import datetime as dt
import sqlite3
from pathlib import Path

from openpyxl import Workbook

from ingest.migrate_from_excel import DAILY_COLUMNS, MARKET_VIBE_COLUMNS
from ops.init_db import main as init_db

SECTOR_NAMES = ["Nifty Bank", "Nifty IT"]


def day(n: int) -> dt.datetime:
    """n-th weekday-ish date starting Mon 2026-01-05 (weekends skipped)."""
    d, count = dt.date(2026, 1, 5), 0
    while True:
        if d.weekday() < 5:
            if count == n:
                return dt.datetime.combine(d, dt.time())
            count += 1
        d += dt.timedelta(days=1)


def daily_row(n, **overrides):
    row = {"date": day(n), "nifty_open": 100.0 + n, "nifty_high": 102.0 + n,
           "nifty_low": 99.0 + n, "nifty_close": 101.0 + n}
    row.update(overrides)
    return row


def make_workbook(path: Path, daily=(), sectors=(), events=(), lists=SECTOR_NAMES,
                  reverse_daily_columns=False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "DailyData"
    cols = ["date"] + DAILY_COLUMNS
    if reverse_daily_columns:
        cols = list(reversed(cols))
    ws.append(cols)
    for rec in daily:
        ws.append([rec.get(c) for c in cols])

    ws = wb.create_sheet("Market_Vibe")
    ws.append(["Date "] + MARKET_VIBE_COLUMNS)

    ws = wb.create_sheet("sectors")
    ws.append(["sector_id", "date", "sector_name", "performance", "sector_return_pct"])
    for r in sectors:
        ws.append(list(r))

    ws = wb.create_sheet("Events_page ")  # trailing space, as in the real workbook
    ws.append(["event_id", "date", "event_type", "event_category",
               "event_description", "impact_direction", "impact_magnitude"])
    for r in events:
        ws.append(list(r))

    wb.create_sheet("Learnings").append(
        ["learning_id", "date", "lesson_priority", "lesson_text", "lesson_category", "actionable"])
    wb.create_sheet("Metadata").append(["key", "value", "comment"])
    ws = wb.create_sheet("Lists")  # NO header row: column A is the data
    for name in lists:
        ws.append([name])
    wb.save(path)
    return path


def make_db(path: Path) -> Path:
    init_db(path)
    return path


def query(db_path, sql, params=()):
    conn = sqlite3.connect(db_path)
    try:
        return conn.execute(sql, params).fetchall()
    finally:
        conn.close()
