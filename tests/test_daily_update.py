"""
Updater tests (converted from the old print-based harness to real pytest).

Every test uses a temporary workbook and a MockDataSource: no network, and
the real workbook is never opened for writing.
"""

import datetime as dt
import shutil
from zoneinfo import ZoneInfo

import openpyxl
import pytest

import config
import excel_ops
import holiday_state
import workbook_builder as twb
from daily_update import Status, run_update
from data_source import DayBar, MockDataSource

IST = ZoneInfo("Asia/Kolkata")
TODAY = dt.date(2026, 8, 26)
AFTER_CLOSE = dt.datetime(2026, 8, 26, 18, 0, tzinfo=IST)
BEFORE_CLOSE = dt.datetime(2026, 8, 26, 11, 0, tzinfo=IST)

OK_PLAN = {
    "nifty": {"mode": "ok", "bar": DayBar(24175, 24334, 24115, 24334), "prev_close": 24219, "has_bars_after": True},
    "vix": {"mode": "ok", "bar": DayBar(11.53, 11.62, 9.45, 11.08), "prev_close": 11.2, "has_bars_after": True},
    "brent": {"mode": "ok", "bar": DayBar(90, 91, 89, 90.5), "prev_close": 91.35, "has_bars_after": True},
    "usdinr": {"mode": "ok", "bar": DayBar(95.3, 95.5, 95.2, 95.4), "prev_close": 95.7, "has_bars_after": True},
    "us10y": {"mode": "ok", "bar": DayBar(46, 47, 45, 46.5), "prev_close": 47.1, "has_bars_after": True},
}
PREV = {"date": dt.date(2026, 8, 25), "nifty_close": 24219, "advance_count": 23, "decline_count": 26}


@pytest.fixture
def xl(tmp_path, monkeypatch):
    """Isolated state: private holiday file and backup dir; returns workbook path."""
    monkeypatch.setattr(config, "CONFIRMED_HOLIDAYS_FILE", tmp_path / "confirmed_holidays.json")
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path)
    return tmp_path / "Market_Database.xlsx"


def run(xl, plan=OK_PLAN, **kw):
    kw.setdefault("today", TODAY)
    kw.setdefault("now", AFTER_CLOSE)
    return run_update(MockDataSource(plan), excel_path=xl, **kw)


def sheet(xl):
    return openpyxl.load_workbook(xl)[config.SHEET_NAME]


def dated_rows(xl):
    ws = sheet(xl)
    return sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value)


def test_normal_successful_day(xl):
    twb.build_workbook(xl, [PREV])
    assert run(xl).status == Status.SUCCESS_CATCHUP_COMPLETE
    assert sheet(xl).cell(row=3, column=5).value == 24334.0


def test_run_twice_same_day(xl):
    twb.build_workbook(xl, [PREV])
    assert run(xl).status == Status.SUCCESS_CATCHUP_COMPLETE
    assert run(xl).status == Status.NO_ACTION_ALREADY_COMPLETE
    assert dated_rows(xl) == 2


def test_weekend_skipped(xl):
    twb.build_workbook(xl, [{"date": dt.date(2026, 8, 21), "nifty_close": 24252}])
    assert run(xl, explicit_date=dt.date(2026, 8, 22)).status == Status.NO_ACTION_WEEKEND


def test_unconfirmed_missing_day_creates_no_row_and_no_holiday_record(xl):
    """A weekday with no Nifty bar but later bars present is skipped, never invented.
    (By design it is NOT written to confirmed_holidays.json - add holidays by hand.)"""
    twb.build_workbook(xl, [{"date": dt.date(2026, 8, 21), "nifty_close": 24252}])
    plan = {"nifty": {"mode": "no_data", "has_bars_after": True}}
    assert run(xl, plan).status == Status.NO_ACTION_NOTHING_TO_DO
    assert dated_rows(xl) == 1
    assert dt.date(2026, 8, 24) not in holiday_state.load(config.CONFIRMED_HOLIDAYS_FILE)


def test_confirmed_holiday_is_skipped_without_network(xl):
    twb.build_workbook(xl, [{"date": dt.date(2026, 8, 21), "nifty_close": 24252}])
    holiday_state.add(config.CONFIRMED_HOLIDAYS_FILE, dt.date(2026, 8, 24))
    result = run(xl, {"nifty": {"mode": "network_error"}}, today=dt.date(2026, 8, 24))
    assert result.status == Status.NO_ACTION_ALREADY_COMPLETE


def test_todays_data_not_yet_available(xl):
    twb.build_workbook(xl, [PREV])
    plan = {"nifty": {"mode": "no_data", "has_bars_after": False}}
    assert run(xl, plan).status == Status.NO_ACTION_DATA_NOT_YET_AVAILABLE
    assert dated_rows(xl) == 1


def test_network_failure_leaves_workbook_untouched(xl):
    twb.build_workbook(xl, [PREV])
    original = xl.read_bytes()
    assert run(xl, {"nifty": {"mode": "network_error"}}).status == Status.ERROR_NETWORK
    assert xl.read_bytes() == original


def test_missing_workbook_reported(xl):
    assert run(xl).status == Status.ERROR_WORKBOOK_PROBLEM


def test_corrupt_workbook_reported_not_unexpected(xl):
    xl.write_bytes(b"this is not a real xlsx file")
    assert run(xl).status == Status.ERROR_WORKBOOK_PROBLEM


def test_invalid_ohlc_rejected_and_not_written(xl):
    twb.build_workbook(xl, [PREV])
    original = xl.read_bytes()
    bad = DayBar(open=24175, high=24000, low=24115, close=24334)  # high < close
    plan = {"nifty": {"mode": "ok", "bar": bad, "prev_close": 24219, "has_bars_after": True}}
    assert run(xl, plan).status == Status.ERROR_INVALID_DATA
    assert xl.read_bytes() == original


def test_incomplete_row_is_repaired(xl):
    twb.build_workbook(xl, [PREV, {"date": dt.date(2026, 8, 26)}])
    assert run(xl).status == Status.SUCCESS_CATCHUP_COMPLETE
    assert sheet(xl).cell(row=3, column=5).value == 24334.0


def test_manual_fields_preserved(xl):
    twb.build_workbook(xl, [PREV, {"date": dt.date(2026, 8, 26), "emotional_state": "MY OWN NOTE"}])
    run(xl)
    ws = sheet(xl)
    col = excel_ops.get_header_map(ws)["emotional_state"]
    assert ws.cell(row=3, column=col).value == "MY OWN NOTE"


def test_duplicate_dates_refused(xl):
    twb.build_workbook(xl, [PREV, {"date": dt.date(2026, 8, 25)}])
    assert run(xl).status == Status.ERROR_DUPLICATE_DATES


def test_future_date_refused(xl):
    twb.build_workbook(xl, [PREV])
    assert run(xl, explicit_date=dt.date(2026, 8, 27)).status == Status.NO_ACTION_FUTURE_DATE


def test_failing_secondary_ticker_does_not_block_nifty(xl):
    twb.build_workbook(xl, [PREV])
    plan = dict(OK_PLAN, vix={"mode": "network_error"}, brent={"mode": "malformed"})
    assert run(xl, plan).status == Status.SUCCESS_CATCHUP_COMPLETE
    ws = sheet(xl)
    h = excel_ops.get_header_map(ws)
    assert ws.cell(row=3, column=h["nifty_close"]).value == 24334.0
    assert ws.cell(row=3, column=h["vix_close"]).value in (None, "")


def test_never_overwrites_existing_value(xl):
    twb.build_workbook(xl, [PREV, {"date": dt.date(2026, 8, 26), "nifty_open": 99999}])
    run(xl)
    ws = sheet(xl)
    assert ws.cell(row=3, column=excel_ops.get_header_map(ws)["nifty_open"]).value == 99999


def test_empty_sheet_without_seed_date_refused(xl):
    twb.build_workbook(xl, [])
    assert run(xl).status == Status.ERROR_WORKBOOK_PROBLEM


def test_path_independent_of_cwd(xl, tmp_path, monkeypatch):
    twb.build_workbook(xl, [PREV])
    other = tmp_path / "elsewhere"
    other.mkdir()
    monkeypatch.chdir(other)
    assert run(xl).status == Status.SUCCESS_CATCHUP_COMPLETE


@pytest.mark.skipif(not config.EXCEL_PATH.exists(), reason="live workbook not present")
def test_real_workbook_structure_matches_config(tmp_path):
    copy = tmp_path / "copy.xlsx"
    shutil.copy2(config.EXCEL_PATH, copy)  # works on a COPY only
    wb = excel_ops.open_workbook(copy)
    ws, _ = excel_ops.get_sheet_and_table(wb)
    headers = excel_ops.get_header_map(ws)
    excel_ops.check_required_headers(headers)
    assert not excel_ops.find_duplicate_dates(ws, headers)


# --- market-close guard (prevents partial intraday bars being locked in) ---------

def test_today_refused_before_market_close(xl):
    twb.build_workbook(xl, [PREV])
    original = xl.read_bytes()
    result = run(xl, now=BEFORE_CLOSE)
    assert result.status == Status.NO_ACTION_DATA_NOT_YET_AVAILABLE
    assert xl.read_bytes() == original


def test_explicit_today_refused_before_close(xl):
    twb.build_workbook(xl, [PREV])
    assert run(xl, now=BEFORE_CLOSE, explicit_date=TODAY).status == Status.NO_ACTION_DATA_NOT_YET_AVAILABLE


def test_catchup_fills_past_days_but_not_todays_partial_bar(xl):
    twb.build_workbook(xl, [{"date": dt.date(2026, 8, 24), "nifty_close": 24000, "advance_count": 20, "decline_count": 20}])
    assert run(xl, now=BEFORE_CLOSE).status == Status.SUCCESS_CATCHUP_COMPLETE
    ws = sheet(xl)
    dates = [ws.cell(row=r, column=1).value.date() for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert dates == [dt.date(2026, 8, 24), dt.date(2026, 8, 25)]  # 26th (today) not written


def test_today_accepted_after_close(xl):
    twb.build_workbook(xl, [PREV])
    assert run(xl, now=dt.datetime(2026, 8, 26, 15, 45, tzinfo=IST)).status == Status.SUCCESS_CATCHUP_COMPLETE
