"""Migration tests: temp workbook + temp DB only. The live files are never used."""

import sqlite3

import pytest

from helpers import daily_row, day, make_db, make_workbook, query
from ingest.migrate_from_excel import main as migrate


def run(tmp_path, **wb_kwargs):
    wb = make_workbook(tmp_path / "wb.xlsx", **wb_kwargs)
    db = tmp_path / "t.db"
    if not db.exists():
        make_db(db)
    migrate(db_path=db, wb_path=wb)
    return db


def test_fresh_rebuild_loads_sector_list_before_sectors(tmp_path):
    sectors = [(1, day(0), "Nifty Bank", "TOP_GAINER", 1.2), (2, day(0), "Nifty IT", "TOP_LOSER", -0.8)]
    db = run(tmp_path, daily=[daily_row(0)], sectors=sectors)
    assert query(db, "SELECT COUNT(*) FROM sector_list") == [(2,)]
    assert query(db, "SELECT COUNT(*) FROM sectors") == [(2,)]
    assert query(db, "SELECT COUNT(*) FROM data_quality_issues") == [(0,)]


def test_blank_manual_fields_are_filled_on_rerun_but_values_are_never_overwritten(tmp_path):
    db = run(tmp_path, daily=[daily_row(0)])  # manual fields blank at first import
    assert query(db, "SELECT fii_equity_cr FROM daily_data") == [(None,)]

    # You later fill FII in Excel AND (wrongly) the Excel nifty_close changes.
    run(tmp_path, daily=[daily_row(0, fii_equity_cr=-1500.0, nifty_close=555.0)])
    fii, close = query(db, "SELECT fii_equity_cr, nifty_close FROM daily_data")[0]
    assert fii == -1500.0   # blank filled in
    assert close == 101.0   # existing value kept


def test_ratio_is_computed_when_excel_has_no_cached_formula_result(tmp_path):
    db = run(tmp_path, daily=[daily_row(0, advance_count=17, decline_count=32)])
    assert query(db, "SELECT advance_decline_ratio FROM daily_data")[0][0] == pytest.approx(17 / 32)


def test_zero_declines_does_not_crash_and_leaves_ratio_null(tmp_path):
    db = run(tmp_path, daily=[daily_row(0, advance_count=10, decline_count=0)])
    assert query(db, "SELECT advance_decline_ratio FROM daily_data") == [(None,)]


def test_columns_are_mapped_by_name_not_position(tmp_path):
    db = run(tmp_path, daily=[daily_row(0, fii_equity_cr=5.0)], reverse_daily_columns=True)
    assert query(db, "SELECT nifty_close, fii_equity_cr FROM daily_data") == [(101.0, 5.0)]


def test_rerun_is_idempotent(tmp_path):
    rows = dict(daily=[daily_row(0), daily_row(1)])
    db = run(tmp_path, **rows)
    run(tmp_path, **rows)
    assert query(db, "SELECT COUNT(*) FROM daily_data") == [(2,)]


def test_blank_template_rows_are_not_reported_as_issues(tmp_path):
    db = run(tmp_path, daily=[daily_row(0), {c: None for c in daily_row(0)}])
    assert query(db, "SELECT COUNT(*) FROM data_quality_issues") == [(0,)]


def test_missing_required_column_fails_loudly(tmp_path):
    wb = make_workbook(tmp_path / "wb.xlsx", daily=[daily_row(0)])
    import openpyxl
    book = openpyxl.load_workbook(wb)
    book["DailyData"].cell(row=1, column=5).value = "renamed_col"
    book.save(wb)
    db = make_db(tmp_path / "t.db")
    with pytest.raises(ValueError, match="missing required columns"):
        migrate(db_path=db, wb_path=wb)
    assert query(db, "SELECT COUNT(*) FROM daily_data") == [(0,)]  # rolled back / nothing written


def test_unknown_column_in_generic_sheet_is_rejected(tmp_path):
    wb = make_workbook(tmp_path / "wb.xlsx", daily=[daily_row(0)])
    import openpyxl
    book = openpyxl.load_workbook(wb)
    book["sectors"].cell(row=1, column=3).value = "sector_name); DROP TABLE daily_data;--"
    book["sectors"].append([1, day(0), "Nifty Bank", "TOP_GAINER", 1.0])
    book.save(wb)
    db = make_db(tmp_path / "t.db")
    with pytest.raises(ValueError, match="not in table"):
        migrate(db_path=db, wb_path=wb)
    assert query(db, "SELECT name FROM sqlite_master WHERE name='daily_data'")  # still there
