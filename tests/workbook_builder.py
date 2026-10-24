"""
test_workbook_builder.py
-------------------------
Builds small synthetic .xlsx files that mirror the real DailyData sheet's
structure (same headers, same Excel Table, same formula column) so the
test suite can exercise table-extension / formula-preservation / manual-
field-preservation logic without ever touching the user's real workbook.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import List, Dict, Optional

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

import config

ALL_HEADERS = (
    ["date"]
    + list(config.AUTO_FIELD_MAP.keys())
    + ["fii_equity_cr", "dii_equity_cr", "fii_index_futures_cr", "fii_index_options_cr",
       "fii_stock_futures_cr", "pcr_value", "pcr_type", "resistance_imm", "resistance_max",
       "support_imm", "support_max", "advance_count", "decline_count", "advance_decline_ratio"]
    + ["emotional_state", "tomorrow_bias", "prediction_hit"]
)
# de-duplicate while preserving order (nifty_change_pts etc. already in AUTO_FIELD_MAP? no)
def _ordered_unique(seq):
    seen = set()
    out = []
    for x in seq:
        if x not in seen:
            seen.add(x)
            out.append(x)
    return out


FULL_HEADER_LIST = _ordered_unique(
    ["date"]
    + list(config.AUTO_FIELD_MAP.keys())
    + config.DERIVED_COLUMNS
    + list(config.MANUAL_COLUMNS.keys())
    + ["advance_decline_ratio"]
)


def build_workbook(dest_path: Path, existing_rows: Optional[List[Dict]] = None) -> None:
    """
    existing_rows: list of {"date": date, "nifty_close": 24000 or None, ...
    any other header: value}. Rows are written in the given order starting
    at sheet row 2.
    """
    existing_rows = existing_rows or []

    wb = Workbook()
    ws = wb.active
    ws.title = config.SHEET_NAME

    for col_idx, header in enumerate(FULL_HEADER_LIST, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_offset, row_data in enumerate(existing_rows, start=2):
        for header, value in row_data.items():
            col_idx = FULL_HEADER_LIST.index(header) + 1
            if header == "date" and isinstance(value, dt.date) and not isinstance(value, dt.datetime):
                value = dt.datetime.combine(value, dt.time())
            ws.cell(row=row_offset, column=col_idx, value=value)
        adr_col = FULL_HEADER_LIST.index("advance_decline_ratio") + 1
        if ws.cell(row=row_offset, column=adr_col).value in (None, ""):
            ws.cell(row=row_offset, column=adr_col).value = (
                f"={config.TABLE_NAME}[[#This Row],[advance_count]]/"
                f"{config.TABLE_NAME}[[#This Row],[decline_count]]"
            )

    last_row = 1 + len(existing_rows)
    last_col_letter = ws.cell(row=1, column=len(FULL_HEADER_LIST)).coordinate.rstrip("0123456789")
    ref = f"A1:{last_col_letter}{last_row}"

    table = Table(displayName=config.TABLE_NAME, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium7", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)

    wb.save(dest_path)
