"""
ops/validate_excel.py
---------------------
PRE-MIGRATION check of the workbook. Read-only: never edits the workbook or
the database. Run it BEFORE ingest/migrate_from_excel.py.

Exit code: 0 = safe to migrate, 1 = at least one blocking problem.

Blocking (the import would reject, mangle or silently skip the data):
    - required DailyData headers missing
    - a date cell that to_iso_date() cannot read (e.g. '31-08-2026', 'TBD', 46265)
    - the same date on two rows
    - text in a numeric cell (e.g. 'TBD', or a space)
    - OHLC rows that are impossible (high < low, open/close outside the range)
    - confirmed_holidays.json entries that cannot be read

Warnings (the import copes, but you should know):
    - date cells that carry a time of day
    - dates stored as text instead of real Excel dates (mixed types)
    - weekend dates, or a date that is also listed as a confirmed holiday
    - sheets (other than DailyData) whose date column has unreadable values
"""

import sys
import warnings
from collections import Counter
from datetime import date, datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

import config
from ingest.migrate_from_excel import (
    DAILY_COLUMNS, TEXT_DAILY_COLUMNS, _is_blank_row, _norm_header,
    clean_value, coerce_numeric, get_sheet_names, to_iso_date,
)

OHLC = {"nifty": ("nifty_open", "nifty_high", "nifty_low", "nifty_close"),
        "vix": ("vix_open", "vix_high", "vix_low", "vix_close")}


def _holidays():
    sys.path.insert(0, str(config.STATE_DIR))
    try:
        import holiday_state
        return holiday_state.load_with_problems(config.CONFIRMED_HOLIDAYS_FILE)
    finally:
        sys.path.pop(0)


def check_daily_data(ws, holidays):
    errors, warns = [], []
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ["DailyData sheet is empty"], warns
    idx = {_norm_header(h): i for i, h in enumerate(rows[0]) if h is not None}
    missing = [c for c in ["date"] + DAILY_COLUMNS if c not in idx]
    if missing:
        return [f"missing required columns: {missing}"], warns

    seen, kinds = {}, Counter()
    for n, row in enumerate(rows[1:], start=2):
        if _is_blank_row(row):
            continue
        raw = row[idx["date"]]
        iso = to_iso_date(raw)
        if iso is None:
            errors.append(f"row {n}: unreadable date {raw!r} (use a real Excel date or YYYY-MM-DD)")
            continue
        kinds["text" if isinstance(raw, str) else "date"] += 1
        if isinstance(raw, datetime) and (raw.hour or raw.minute or raw.second):
            warns.append(f"row {n}: date {iso} carries a time of day ({raw:%H:%M:%S}); the time is dropped on import")
        d = date.fromisoformat(iso)
        if d.weekday() >= 5:
            warns.append(f"row {n}: {iso} is a {d:%A}")
        if d in holidays:
            warns.append(f"row {n}: {iso} has data but is listed in confirmed_holidays.json")
        if iso in seen:
            errors.append(f"row {n}: duplicate date {iso} (first seen on row {seen[iso]})")
        seen.setdefault(iso, n)

        for col in DAILY_COLUMNS:
            if col in TEXT_DAILY_COLUMNS:
                continue
            _, err = coerce_numeric(clean_value(row[idx[col]]))
            if err:
                errors.append(f"row {n} ({iso}) {col}: {err}")

        for name, (o, h, l, c) in OHLC.items():
            vals = [clean_value(row[idx[k]]) for k in (o, h, l, c)]
            if all(isinstance(v, (int, float)) for v in vals):
                vo, vh, vl, vc = vals
                if vh < vl or vh < max(vo, vc) or vl > min(vo, vc):
                    errors.append(f"row {n} ({iso}) {name}: impossible OHLC {vals}")

    if len(kinds) > 1:
        warns.append(f"date column mixes {kinds['date']} real dates with {kinds['text']} text dates")
    return errors, warns


def check_other_sheet(ws, name):
    """Returns (errors, warnings) for sectors/events/learnings/market_vibe."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    heads = [_norm_header(h) for h in rows[0] if h is not None]
    if "date" not in heads:
        return [], []
    di, errs, warns = heads.index("date"), [], []
    numeric = [(heads.index(c), c) for c in ("sector_return_pct",) if c in heads]
    for n, row in enumerate(rows[1:], start=2):
        if _is_blank_row(row):
            continue
        if to_iso_date(row[di]) is None:
            warns.append(f"{name} row {n}: unreadable date {row[di]!r}")
        for i, col in numeric:
            _, err = coerce_numeric(clean_value(row[i]))
            if err:
                errs.append(f"{name} row {n} {col}: {err}")
    return errs, warns


def main(wb_path=None) -> int:
    wb_path = Path(wb_path or config.EXCEL_PATH)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # openpyxl: unsupported Excel extensions
        wb = load_workbook(wb_path, read_only=True, data_only=True)
    sheets = get_sheet_names(wb)
    holidays, hproblems = _holidays()

    errors = [f"confirmed_holidays.json: {p}" for p in hproblems if "IGNORED" in p or "could not" in p]
    warns = [f"confirmed_holidays.json: {p}" for p in hproblems if p not in errors and "accepted" in p]

    if sheets["daily_data"] not in wb.sheetnames:
        errors.append("DailyData sheet not found")
    else:
        e, w = check_daily_data(wb[sheets["daily_data"]], holidays)
        errors += [f"DailyData {m}" for m in e]
        warns += [f"DailyData {m}" for m in w]
    for key in ("sectors", "events", "learnings", "market_vibe"):
        name = sheets[key]
        if name in wb.sheetnames:
            e, w = check_other_sheet(wb[name], name.strip())
            errors += e
            warns += w
    wb.close()

    print("=" * 60 + f"\nEXCEL PRE-MIGRATION CHECK: {wb_path.name}\n" + "=" * 60)
    for m in warns:
        print(f"  [WARN] {m}")
    for m in errors:
        print(f"  [PROBLEM] {m}")
    if errors:
        print(f"\n  [FAIL] {len(errors)} blocking problem(s). Fix the workbook, then re-run.")
        return 1
    print(f"\n  [OK] Safe to migrate ({len(warns)} warning(s)).")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else None))
