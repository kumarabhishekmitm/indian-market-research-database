"""
ops/apply_excel_formatting.py
-----------------------------
Self-healing formatting engine enforcing Data Validation lists, auto-calculating 
derived metrics only for active rows, and expanding column widths dynamically.
"""

import sys
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import config

EXCEL_PATH = config.EXCEL_PATH

SECTOR_INDEX_LIST = [
    "Nifty Auto", "Nifty Bank", "Nifty Consumer Durables", "Nifty Energy",
    "Nifty Financial Services", "Nifty Financial Services 25/50", "Nifty FMCG",
    "Nifty Healthcare", "Nifty IT", "Nifty Media", "Nifty Metal",
    "Nifty MidSmall Healthcare", "Nifty Oil & Gas", "Nifty Pharma",
    "Nifty Private Bank", "Nifty PSU Bank", "Nifty Realty",
]


def get_header_map(ws):
    headers = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val is not None:
            headers[str(val).strip()] = col
    return headers


def apply_formatting_and_validations(excel_path: Path):
    wb = openpyxl.load_workbook(excel_path)
    
    # Normalize sheet name if trailing space exists
    if "Lists " in wb.sheetnames:
        wb["Lists "].title = "Lists"

    max_rows = 500

    # 1. Market_Vibe Validations
    if "Market_Vibe" in wb.sheetnames:
        ws = wb["Market_Vibe"]
        headers = get_header_map(ws)
        if "Asian market" in headers:
            col = get_column_letter(headers["Asian market"])
            dv = DataValidation(type="list", formula1='"Positive,Negative,Neutral"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max_rows}")
        if "Market_Vibe" in headers:
            col = get_column_letter(headers["Market_Vibe"])
            dv = DataValidation(type="list", formula1='"Risk ON,Risk OFF"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max_rows}")

    # 2. DailyData Validations & Formulas
    if "DailyData" in wb.sheetnames:
        ws = wb["DailyData"]
        ws.data_validations.dataValidation.clear()

        headers = get_header_map(ws)
        rules = {
            "pcr_type": '"WEEKLY,MONTHLY"',
            "emotional_state": '"CALM,CONFIDENT,OPTIMISTIC,NEUTRAL,UNCERTAIN,ANXIOUS,FEARFUL,GREEDY,FRUSTRATED,OVERCONFIDENT"',
            "tomorrow_bias": '"STRONG BULLISH,BULLISH,NEUTRAL,BEARISH,STRONG BEARISH"',
            "prediction_hit": '"0,1"',
        }
        for field_name, formula in rules.items():
            if field_name in headers:
                col = get_column_letter(headers[field_name])
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{max_rows}")

        # Auto-calculate Advance/Decline Ratio ONLY if inputs are present
        if (
            "advance_count" in headers
            and "decline_count" in headers
            and "advance_decline_ratio" in headers
        ):
            adv_col = get_column_letter(headers["advance_count"])
            dec_col = get_column_letter(headers["decline_count"])
            ratio_col = get_column_letter(headers["advance_decline_ratio"])

            for r in range(2, max_rows + 1):
                ws[f"{ratio_col}{r}"] = f'=IF(OR({adv_col}{r}="", {dec_col}{r}=""), "", IF({dec_col}{r}=0, 0, {adv_col}{r}/{dec_col}{r}))'

    # 3. Sectors Validations
    if "sectors" in wb.sheetnames:
        ws = wb["sectors"]
        ws.data_validations.dataValidation.clear()
        headers = get_header_map(ws)
        if "sector_name" in headers:
            col = get_column_letter(headers["sector_name"])
            dv = DataValidation(type="list", formula1="'Lists'!$A$1:$A$17", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max_rows}")
        if "performance" in headers:
            col = get_column_letter(headers["performance"])
            dv = DataValidation(type="list", formula1='"TOP_GAINER,TOP_LOSER,NEUTRAL"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max_rows}")

    # 4. Dynamic Column Auto-Fit
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(excel_path)
    print(f"[SUCCESS] Workbook formatted, validated & smart formulas injected: {excel_path}")


if __name__ == "__main__":
    apply_formatting_and_validations(EXCEL_PATH)