"""
ingest/migrate_from_excel.py
----------------------------
Excel -> SQLite ingestion.

Behaviour (changed from the original version):
  * Importing this module has NO side effects (the workbook is only opened
    inside main()).
  * main(db_path, wb_path) takes explicit paths, so tests can use temporary
    copies instead of the live files.
  * daily_data / market_vibe are UPSERTED in "fill the blanks" mode: an
    existing date keeps every value it already has, but any column that is
    NULL in SQLite and filled in Excel is filled in. Existing values are
    never overwritten.
  * advance_decline_ratio is computed from the counts when Excel has no
    cached formula result (openpyxl saves strip cached formula values).
  * daily_data columns are mapped by HEADER NAME, not by position.
  * sector_list is loaded from the 'Lists' sheet BEFORE sectors, so a fresh
    database can be rebuilt from the workbook.
"""

import re
import sqlite3
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import config

DAILY_COLUMNS = [
    "nifty_open", "nifty_high", "nifty_low", "nifty_close", "nifty_change_pts",
    "nifty_change_pct", "vix_open", "vix_high", "vix_low", "vix_close",
    "vix_change_pct", "fii_equity_cr", "dii_equity_cr", "fii_index_futures_cr",
    "fii_index_options_cr", "fii_stock_futures_cr", "pcr_value", "pcr_type",
    "resistance_imm", "resistance_max", "support_imm", "support_max",
    "advance_count", "decline_count", "advance_decline_ratio", "nifty_gap_pts",
    "emotional_state", "tomorrow_bias", "prediction_hit", "brent_crude",
    "usd_inr", "us_10y_yield",
]

# Market_Vibe headers in the workbook are not clean ('Date ', ' Brent_Crude'),
# so this sheet is still mapped by position: date first, then these.
MARKET_VIBE_COLUMNS = [
    "us_sp500_pct", "us_nasdaq_pct", "gift_nifty_pct", "brent_crude",
    "usd_inr", "us_10y_yield", "asian_market", "market_vibe_label",
    "confidence_score",
]


def get_sheet_names(wb):
    """Resolve logical sheet keys to real sheet names (tolerates stray spaces)."""
    names = wb.sheetnames

    def find(pref):
        return next((n for n in names if n.strip().lower() == pref.lower()), pref)

    return {
        "lists": find("Lists"),
        "daily_data": find("DailyData"),
        "market_vibe": find("Market_Vibe"),
        "sectors": find("sectors"),
        "events": find("Events_page"),
        "learnings": find("Learnings"),
        "metadata": find("Metadata"),
    }


def log_issue(conn, source_sheet, row_number, description, raw_row):
    conn.execute(
        """
        INSERT OR IGNORE INTO data_quality_issues
            (source_sheet, source_row_number, issue_description, raw_row_data, logged_at)
        VALUES (?, ?, ?, ?, ?)
        """,
        (source_sheet, row_number, description, str(raw_row),
         datetime.now(timezone.utc).isoformat()),
    )


_LOOSE_YMD = re.compile(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:[ T].*)?$")


def to_iso_date(value):
    """Return 'YYYY-MM-DD' or None.

    Accepts datetime/date objects and year-first text ('2026-08-31',
    '2026-8-31', '2026/08/31', optionally followed by a time). Day-first text
    ('31-08-2026') and bare numbers are deliberately NOT guessed - they are
    ambiguous, so they are returned as None and logged for manual review.
    """
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        text = value.strip()
        m = _LOOSE_YMD.match(text)
        if m:
            try:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
            except ValueError:
                return None
        try:
            return datetime.fromisoformat(text).strftime("%Y-%m-%d")
        except ValueError:
            return None
    return None


TEXT_DAILY_COLUMNS = {"pcr_type", "emotional_state", "tomorrow_bias"}

# Numeric columns of the generic (insert-if-new) sheets: table -> columns.
NUMERIC_GENERIC_COLUMNS = {"sectors": {"sector_return_pct"}}


def coerce_numeric(value):
    """(number_or_None, error_or_None) for a cell that must be numeric."""
    if value is None or isinstance(value, (int, float)) and not isinstance(value, bool):
        return value, None
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", "")), None
        except ValueError:
            pass
    return None, f"expected a number, got {value!r}"


def clean_value(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, str) and value.strip().lower() in config.NA_VALUES:
        return None
    return value


def _is_blank_row(row):
    return all(c is None or (isinstance(c, float) and pd.isna(c)) for c in row)


def _norm_header(h):
    return str(h).strip().lower().replace(" ", "_")


def _upsert_fill_blanks(conn, table, key_col, key_val, values):
    """Insert a row, or fill only the NULL columns of the existing one.

    values: dict of column -> value (key column excluded).
    Returns "inserted", "updated" (>=1 blank filled) or "unchanged".
    """
    existing = conn.execute(
        f"SELECT * FROM {table} WHERE {key_col} = ?", (key_val,)
    ).fetchone()
    cols = [key_col] + list(values)
    placeholders = ",".join("?" for _ in cols)
    assignments = ",".join(
        f"{c}=COALESCE({table}.{c}, excluded.{c})" for c in values
    )
    conn.execute(
        f"INSERT INTO {table} ({','.join(cols)}) VALUES ({placeholders}) "
        f"ON CONFLICT({key_col}) DO UPDATE SET {assignments}",
        [key_val] + list(values.values()),
    )
    if existing is None:
        return "inserted"
    names = [d[0] for d in conn.execute(f"SELECT * FROM {table} LIMIT 0").description]
    old = dict(zip(names, existing))
    filled = any(old[c] is None and v is not None for c, v in values.items())
    return "updated" if filled else "unchanged"


def _new_counts():
    return {"inserted": 0, "updated": 0, "unchanged": 0, "blank": 0, "issues": 0}


def migrate_sector_list(wb, conn, sheets):
    """'Lists' has NO header row: column A is the list of valid sector names."""
    name = sheets["lists"]
    counts = _new_counts()
    if name not in wb.sheetnames:
        return counts
    for (value,) in wb[name].iter_rows(min_col=1, max_col=1, values_only=True):
        value = clean_value(value)
        if value is None:
            counts["blank"] += 1
            continue
        cur = conn.execute(
            "INSERT OR IGNORE INTO sector_list (sector_name) VALUES (?)",
            (str(value).strip(),),
        )
        counts["inserted" if cur.rowcount else "unchanged"] += 1
    return counts


def migrate_daily_data(wb, conn, sheets):
    name = sheets["daily_data"]
    counts = _new_counts()
    if name not in wb.sheetnames:
        return counts

    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return counts
    header_idx = {_norm_header(h): i for i, h in enumerate(rows[0]) if h is not None}
    missing = [c for c in ["date"] + DAILY_COLUMNS if c not in header_idx]
    if missing:
        raise ValueError(f"DailyData sheet is missing required columns: {missing}")

    for row_number, row in enumerate(rows[1:], start=2):
        if _is_blank_row(row):
            counts["blank"] += 1
            continue
        date_str = to_iso_date(row[header_idx["date"]])
        if not date_str:
            log_issue(conn, name, row_number, f"Invalid date: {row[header_idx['date']]!r}", row)
            counts["issues"] += 1
            continue

        values = {c: clean_value(row[header_idx[c]]) for c in DAILY_COLUMNS}

        # Type guard: text in a numeric column (e.g. '' or 'TBD') would be stored
        # as-is by SQLite and later block the fill-the-blanks logic.
        for col in DAILY_COLUMNS:
            if col in TEXT_DAILY_COLUMNS:
                continue
            values[col], err = coerce_numeric(values[col])
            if err:
                log_issue(conn, name, row_number, f"{date_str} {col}: {err}", row)
                counts["issues"] += 1

        # Derived value: don't depend on Excel's cached formula result.
        adv, dec = values["advance_count"], values["decline_count"]
        if values["advance_decline_ratio"] is None and adv is not None and dec:
            values["advance_decline_ratio"] = adv / dec

        counts[_upsert_fill_blanks(conn, "daily_data", "date", date_str, values)] += 1
    return counts


def migrate_market_vibe(wb, conn, sheets):
    name = sheets["market_vibe"]
    counts = _new_counts()
    if name not in wb.sheetnames:
        return counts
    rows = list(wb[name].iter_rows(values_only=True))[1:]
    for row_number, row in enumerate(rows, start=2):
        if _is_blank_row(row):
            counts["blank"] += 1
            continue
        date_str = to_iso_date(row[0])
        if not date_str:
            log_issue(conn, name, row_number, f"Invalid or missing date in Market_Vibe: {row[0]!r}", row)
            counts["issues"] += 1
            continue
        values = {c: clean_value(row[i]) for i, c in enumerate(MARKET_VIBE_COLUMNS, start=1)}
        try:
            counts[_upsert_fill_blanks(conn, "market_vibe", "date", date_str, values)] += 1
        except sqlite3.Error as e:
            log_issue(conn, name, row_number, f"SQL Insert Error: {e}", row)
            counts["issues"] += 1
    return counts


def migrate_generic_dated_sheet(wb, conn, sheets, sheet_key, table_name):
    """sectors / events / learnings / metadata: insert-if-new, keyed on the
    sheet's own id (or key) column. Existing rows are left untouched."""
    name = sheets.get(sheet_key)
    counts = _new_counts()
    if not name or name not in wb.sheetnames:
        return counts
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return counts

    headers = [_norm_header(h) for h in rows[0] if h is not None]
    table_cols = {r[1] for r in conn.execute(f"PRAGMA table_info({table_name})")}
    unknown = [h for h in headers if h not in table_cols]
    if unknown:  # also prevents header text from ever being spliced into SQL
        raise ValueError(f"Sheet '{name}' has columns not in table '{table_name}': {unknown}")
    date_idx = headers.index("date") if "date" in headers else None

    for row_number, row in enumerate(rows[1:], start=2):
        if _is_blank_row(row):
            counts["blank"] += 1
            continue
        vals = [clean_value(v) for v in row[: len(headers)]]
        for i, h in enumerate(headers):
            if h in NUMERIC_GENERIC_COLUMNS.get(table_name, ()):
                vals[i], err = coerce_numeric(vals[i])
                if err:
                    log_issue(conn, name, row_number, f"{h}: {err}", row)
                    counts["issues"] += 1
        if date_idx is not None:
            date_str = to_iso_date(row[date_idx])
            if not date_str:
                log_issue(conn, name, row_number, f"Invalid or missing date: {row[date_idx]!r}", row)
                counts["issues"] += 1
                continue
            vals[date_idx] = date_str
        try:
            cur = conn.execute(
                f"INSERT OR IGNORE INTO {table_name} ({','.join(headers)}) "
                f"VALUES ({','.join('?' for _ in vals)})",
                vals,
            )
            counts["inserted" if cur.rowcount else "unchanged"] += 1
        except sqlite3.Error as e:
            log_issue(conn, name, row_number, f"SQL Insert Error: {e}", row)
            counts["issues"] += 1
    return counts


def _report(label, c):
    print(
        f"{label}: {c['inserted']} inserted, {c['updated']} filled-in, "
        f"{c['unchanged']} unchanged, {c['blank']} blank rows, {c['issues']} issues"
    )


def main(db_path=None, wb_path=None):
    db_path = Path(db_path or config.DB_PATH)
    wb_path = Path(wb_path or config.EXCEL_PATH)

    wb = load_workbook(wb_path, read_only=True, data_only=True)
    sheets = get_sheet_names(wb)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON;")

    try:
        _report("sector_list", migrate_sector_list(wb, conn, sheets))  # must precede sectors
        _report("daily_data", migrate_daily_data(wb, conn, sheets))
        _report("market_vibe", migrate_market_vibe(wb, conn, sheets))
        for key, table in [("sectors", "sectors"), ("events", "events"),
                           ("learnings", "learnings"), ("metadata", "metadata")]:
            _report(table, migrate_generic_dated_sheet(wb, conn, sheets, key, table))
        conn.commit()
        print("[SUCCESS] All workbook sheets successfully migrated to SQLite.")
    except Exception as e:
        conn.rollback()
        print(f"[FAIL] Rollback triggered during migration: {e}")
        raise
    finally:
        conn.close()
        wb.close()


if __name__ == "__main__":
    main()
