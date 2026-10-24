"""
Step 3 (updated): Inspect headers, sample data, and distinguish
real data rows from "phantom" formatted-but-empty rows.
Still read-only.
"""

from openpyxl import load_workbook

WORKBOOK_PATH = "../data/Market_Database.xlsx"

wb = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print("=" * 70)
    print(f"SHEET: {sheet_name!r}")
    print("=" * 70)

    all_rows = list(ws.iter_rows(values_only=True))
    headers = all_rows[0]
    data_rows = all_rows[1:]

    # A row is "blank" if every single cell in it is None
    blank_rows = [r for r in data_rows if all(cell is None for cell in r)]
    real_rows = [r for r in data_rows if not all(cell is None for cell in r)]

    print("Headers (repr shows hidden spaces):")
    for h in headers:
        print(f"  {h!r}")

    print(f"\nTotal rows below header: {len(data_rows)}")
    print(f"  -> Completely blank rows: {len(blank_rows)}")
    print(f"  -> Rows with real data:   {len(real_rows)}")

    print("\nFirst 2 real data rows:")
    for row in real_rows[:2]:
        print(f"  {row}")

    print("\nLast 2 real data rows:")
    for row in real_rows[-2:]:
        print(f"  {row}")

    print()

wb.close()